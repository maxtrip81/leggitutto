"""Documenti da ufficio: Word, OpenDocument, RTF, PowerPoint, fogli di calcolo."""

from __future__ import annotations

import csv
import io
import re
import tempfile
import zipfile
from pathlib import Path

from lxml import etree

from . import external
from .i18n import _
from .model import ConversionError, Context, Extracted, Para, ProtectedError
from .settings import Options


# ---------------------------------------------------------------------------
# Tabelle: per chi usa la sintesi vocale "Intestazione: valore" è molto più chiaro
# di colonne allineate con spazi.

def table_to_paras(rows: list[list[str]]) -> list[Para]:
    rows = [[re.sub(r"\s+", " ", c or "").strip() for c in row] for row in rows]
    rows = [row for row in rows if any(row)]
    if not rows:
        return []
    width = max(len(r) for r in rows)
    if width == 1:
        return [Para(r[0]) for r in rows if r[0]]
    header = rows[0]
    use_header = len(rows) > 1 and sum(1 for c in header if c) >= max(2, (width + 1) // 2) and len(set(header)) == len(header)
    out = []
    if use_header:
        for row in rows[1:]:
            cells = [f"{header[i] if i < len(header) and header[i] else _('Colonna {n}').format(n=i + 1)}: {c}" for i, c in enumerate(row) if c]
            if cells:
                out.append(Para("; ".join(cells) + "."))
    else:
        for row in rows:
            cells = [c for c in row if c]
            if cells:
                out.append(Para("; ".join(cells) + "."))
    return out


# ---------------------------------------------------------------------------
# File Office cifrati (password di apertura)

def decrypt_office(path: Path, ctx: Context) -> io.BytesIO | None:
    """Se il file Office è cifrato chiede la password e restituisce il file in chiaro."""
    try:
        import msoffcrypto
    except ImportError:
        return None
    with open(path, "rb") as fh:
        head = fh.read(8)
    if head != b"\xd0\xcf\x11\xe0\xa1\xb1\x1a\xe1":
        return None
    with open(path, "rb") as fh:
        try:
            office = msoffcrypto.OfficeFile(fh)
            if not office.is_encrypted():
                return None
        except Exception:
            return None
        # Excel usa una password predefinita per i file "protetti" ma apribili da chiunque.
        for attempt, password in enumerate(["VelvetSweatshop"] + [None] * 3):
            if password is None:
                password = ctx.ask_password(path.name, attempt - 1)
                if password is None:
                    raise ProtectedError(_("il documento è protetto da password e non è stata inserita la password."))
            try:
                fh.seek(0)
                office = msoffcrypto.OfficeFile(fh)
                office.load_key(password=password)
                out = io.BytesIO()
                office.decrypt(out)
                out.seek(0)
                return out
            except Exception:
                continue
    raise ProtectedError(_("la password inserita non è corretta."))


# ---------------------------------------------------------------------------
# Word .docx

def _docx_paragraph_kind(paragraph) -> str:
    try:
        style = (paragraph.style.name or "").lower()
    except Exception:
        style = ""
    if style.startswith(("heading", "titolo", "title", "intestazione")) and "toc" not in style:
        return "heading"
    if "list" in style or "elenco" in style or paragraph._p.pPr is not None and paragraph._p.pPr.numPr is not None:
        return "list"
    return "text"


def extract_docx(path: Path, opts: Options, ctx: Context, data=None) -> Extracted:
    import docx
    from docx.table import Table
    from docx.text.paragraph import Paragraph

    data = data or decrypt_office(path, ctx)
    try:
        document = docx.Document(data or str(path))
    except Exception as exc:
        raise ConversionError(_("documento Word non leggibile: {exc}").format(exc=exc)) from exc
    paras: list[Para] = []
    w = "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}"

    def walk(container, parent):
        for child in container.iterchildren():
            if child.tag == w + "p":
                p = Paragraph(child, parent)
                text = p.text
                if text.strip():
                    kind = _docx_paragraph_kind(p)
                    for piece in text.split("\n"):
                        if piece.strip():
                            paras.append(Para(piece, kind=kind))
            elif child.tag == w + "tbl":
                table = Table(child, parent)
                rows = []
                for row in table.rows:
                    cells, seen = [], set()
                    for cell in row.cells:
                        if id(cell._tc) in seen:  # celle unite
                            continue
                        seen.add(id(cell._tc))
                        cells.append(" ".join(p.text for p in cell.paragraphs))
                    rows.append(cells)
                paras.extend(table_to_paras(rows))
            elif child.tag == w + "sdt":  # controlli contenuto
                content = child.find(w + "sdtContent")
                if content is not None:
                    walk(content, parent)

    walk(document.element.body, document._body)
    ctx.progress(1.0)
    return Extracted(paras=paras)


# ---------------------------------------------------------------------------
# OpenDocument (.odt, .ods, .odp) senza dipendenze esterne

_ODF = {
    "text": "urn:oasis:names:tc:opendocument:xmlns:text:1.0",
    "table": "urn:oasis:names:tc:opendocument:xmlns:table:1.0",
    "draw": "urn:oasis:names:tc:opendocument:xmlns:drawing:1.0",
    "office": "urn:oasis:names:tc:opendocument:xmlns:office:1.0",
    "presentation": "urn:oasis:names:tc:opendocument:xmlns:presentation:1.0",
}


def _q(prefix, name):
    return "{%s}%s" % (_ODF[prefix], name)


def _odf_inline_text(el, opts: Options) -> str:
    out = [el.text or ""]
    for child in el:
        tag = child.tag
        if tag == _q("text", "note"):
            if not opts.remove_footnotes and not opts.remove_note_markers:
                out.append(" " + _odf_inline_text(child, opts))
        elif tag == _q("text", "s"):
            out.append(" " * int(child.get(_q("text", "c"), "1")))
        elif tag == _q("text", "tab"):
            out.append(" ")
        elif tag == _q("text", "line-break"):
            out.append("\n")
        elif tag in (_q("text", "page-number"), _q("text", "page-count")):
            pass
        elif tag == _q("office", "annotation"):
            pass  # commenti a margine
        else:
            out.append(_odf_inline_text(child, opts))
        out.append(child.tail or "")
    return "".join(out)


def _odf_blocks(el, opts: Options, paras: list[Para], notes: list[Para]):
    for child in el:
        tag = child.tag
        if tag in (_q("text", "p"), _q("text", "h")):
            kind = "heading" if tag == _q("text", "h") else "text"
            text = _odf_inline_text(child, opts)
            if text.strip():
                paras.append(Para(text, kind=kind))
            if not opts.remove_footnotes:
                for note in child.iter(_q("text", "note-body")):
                    _odf_blocks(note, opts, notes, notes)
            # Le cornici (text:p > draw:frame > draw:text-box) contengono altri paragrafi.
            for box in child.iter(_q("draw", "text-box")):
                _odf_blocks(box, opts, paras, notes)
        elif tag == _q("table", "table"):
            rows = []
            for row in child.iter(_q("table", "table-row")):
                cells = []
                for cell in row:
                    if cell.tag == _q("table", "table-cell"):
                        repeat = min(int(cell.get(_q("table", "number-columns-repeated"), "1")), 50)
                        value = " ".join(_odf_inline_text(p, opts) for p in cell.iter(_q("text", "p")))
                        cells.extend([value] * (repeat if value else 1))
                rows.append(cells)
            if child.getparent() is not None and child.getparent().tag == _q("office", "spreadsheet"):
                paras.append(Para(_("Foglio: {name}").format(name=child.get(_q("table", "name"), "")), kind="heading"))
            paras.extend(table_to_paras(rows))
        elif tag == _q("draw", "page"):
            slide_label = _("Diapositiva {n}").format(n="")
            number = len([p for p in paras if p.kind == "heading" and p.text.startswith(slide_label)]) + 1
            paras.append(Para(_("Diapositiva {n}").format(n=number), kind="heading"))
            _odf_blocks(child, opts, paras, notes)
        elif tag == _q("presentation", "notes"):
            if opts.include_slide_notes:
                sub: list[Para] = []
                _odf_blocks(child, opts, sub, notes)
                if sub:
                    paras.append(Para(_("Note del relatore:")))
                    paras.extend(sub)
        elif tag in (_q("text", "table-of-content"), _q("text", "illustration-index")):
            body = child.find(_q("text", "index-body"))
            if body is not None:
                _odf_blocks(body, opts, paras, notes)
        else:
            _odf_blocks(child, opts, paras, notes)


def extract_odf(path: Path, opts: Options, ctx: Context) -> Extracted:
    try:
        with zipfile.ZipFile(path) as zf:
            if "META-INF/manifest.xml" in zf.namelist():
                manifest = zf.read("META-INF/manifest.xml")
                if b"encryption-data" in manifest:
                    raise ProtectedError(_(
                        "il documento OpenDocument è cifrato con password: aprilo con LibreOffice "
                        "e salvalo senza password, poi riprova."))
            root = etree.fromstring(zf.read("content.xml"), etree.XMLParser(huge_tree=True, recover=True))
    except (zipfile.BadZipFile, KeyError) as exc:
        raise ConversionError(_("documento OpenDocument non leggibile: {exc}").format(exc=exc)) from exc
    body = root.find(_q("office", "body"))
    paras: list[Para] = []
    notes: list[Para] = []
    _odf_blocks(body if body is not None else root, opts, paras, notes)
    if notes:
        paras.append(Para(_("Note"), kind="heading"))
        paras.extend(notes)
    return Extracted(paras=paras)


# ---------------------------------------------------------------------------
# RTF

def _strip_rtf_groups(rtf: str, keywords: tuple[str, ...]) -> str:
    """Toglie i gruppi {\\footnote ...} e simili, rispettando le parentesi annidate."""
    pattern = re.compile(r"\{\s*(?:\\\*\s*)?\\(?:%s)\b" % "|".join(keywords))
    out, pos = [], 0
    while True:
        match = pattern.search(rtf, pos)
        if not match:
            out.append(rtf[pos:])
            break
        out.append(rtf[pos:match.start()])
        depth, i = 0, match.start()
        while i < len(rtf):
            ch = rtf[i]
            if ch == "\\":
                i += 2
                continue
            if ch == "{":
                depth += 1
            elif ch == "}":
                depth -= 1
                if depth == 0:
                    break
            i += 1
        pos = i + 1
    return "".join(out)


def extract_rtf(path: Path, opts: Options, ctx: Context) -> Extracted:
    from striprtf.striprtf import rtf_to_text

    raw = path.read_bytes().decode("latin-1")
    groups = ["header", "footer", "headerl", "headerr", "headerf", "footerl", "footerr", "footerf", "annotation", "atnid", "atnauthor"]
    if opts.remove_footnotes:
        groups.append("footnote")
    raw = _strip_rtf_groups(raw, tuple(groups))
    raw = raw.replace("\\chpgn", "")
    text = rtf_to_text(raw, errors="ignore")
    return Extracted(paras=[Para(p) for p in re.split(r"\n", text) if p.strip()])


# ---------------------------------------------------------------------------
# PowerPoint

def extract_pptx(path: Path, opts: Options, ctx: Context, data=None) -> Extracted:
    from pptx import Presentation
    from pptx.enum.shapes import PP_PLACEHOLDER

    data = data or decrypt_office(path, ctx)
    try:
        pres = Presentation(data or str(path))
    except Exception as exc:
        raise ConversionError(_("presentazione non leggibile: {exc}").format(exc=exc)) from exc
    skip = {PP_PLACEHOLDER.SLIDE_NUMBER, PP_PLACEHOLDER.FOOTER, PP_PLACEHOLDER.DATE}
    paras: list[Para] = []

    def shape_paras(shape, out):
        if shape.is_placeholder:
            try:
                if shape.placeholder_format.type in skip:
                    return
            except Exception:
                pass
        if shape.shape_type == 6 and hasattr(shape, "shapes"):  # gruppo
            for sub in sorted(shape.shapes, key=lambda s: ((s.top or 0), (s.left or 0))):
                shape_paras(sub, out)
            return
        if getattr(shape, "has_table", False) and shape.has_table:
            rows = [[cell.text for cell in row.cells] for row in shape.table.rows]
            out.extend(table_to_paras(rows))
            return
        if getattr(shape, "has_text_frame", False) and shape.has_text_frame:
            for p in shape.text_frame.paragraphs:
                text = p.text
                if text.strip():
                    out.append(Para(text, kind="list" if p.level else "text"))

    total = len(pres.slides)
    for n, slide in enumerate(pres.slides, 1):
        ctx.progress(n / max(total, 1), _("Diapositiva {n} di {total}").format(n=n, total=total))
        title_shape = slide.shapes.title
        title = title_shape.text_frame.text.strip() if title_shape is not None and title_shape.has_text_frame else ""
        paras.append(Para(_("Diapositiva {n}").format(n=n) + (f": {title}" if title else ""), kind="heading"))
        shapes = [s for s in slide.shapes if title_shape is None or s.shape_id != title_shape.shape_id]
        for shape in sorted(shapes, key=lambda s: ((s.top or 0), (s.left or 0))):
            shape_paras(shape, paras)
        if opts.include_slide_notes and slide.has_notes_slide:
            notes = slide.notes_slide.notes_text_frame.text.strip() if slide.notes_slide.notes_text_frame else ""
            if notes:
                paras.append(Para(_("Note del relatore:") + " " + notes))
    return Extracted(paras=paras)


# ---------------------------------------------------------------------------
# Fogli di calcolo

def _sheet_paras(name: str, rows, paras: list[Para], multiple: bool):
    rows = [["" if v is None else _cell_str(v) for v in row] for row in rows]
    if multiple:
        paras.append(Para(_("Foglio: {name}").format(name=name), kind="heading"))
    paras.extend(table_to_paras(rows))


def _cell_str(value) -> str:
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    if hasattr(value, "strftime"):
        return value.strftime("%d/%m/%Y %H:%M").replace(" 00:00", "")
    return str(value)


def extract_xlsx(path: Path, opts: Options, ctx: Context, data=None) -> Extracted:
    import openpyxl

    data = data or decrypt_office(path, ctx)
    try:
        wb = openpyxl.load_workbook(data or str(path), read_only=True, data_only=True)
    except Exception as exc:
        raise ConversionError(_("foglio di calcolo non leggibile: {exc}").format(exc=exc)) from exc
    paras: list[Para] = []
    for ws in wb.worksheets:
        _sheet_paras(ws.title, ws.iter_rows(values_only=True), paras, len(wb.worksheets) > 1)
    return Extracted(paras=paras)


def extract_xls(path: Path, opts: Options, ctx: Context) -> Extracted:
    import xlrd

    data = decrypt_office(path, ctx)
    try:
        book = xlrd.open_workbook(file_contents=data.read()) if data else xlrd.open_workbook(str(path))
    except Exception as exc:
        raise ConversionError(_("foglio di calcolo non leggibile: {exc}").format(exc=exc)) from exc
    paras: list[Para] = []
    for sheet in book.sheets():
        rows = [sheet.row_values(r) for r in range(sheet.nrows)]
        _sheet_paras(sheet.name, rows, paras, book.nsheets > 1)
    return Extracted(paras=paras)


def extract_csv(path: Path, opts: Options, ctx: Context) -> Extracted:
    from .extract_web import read_text_file

    text = read_text_file(path)
    try:
        dialect = csv.Sniffer().sniff(text[:5000], delimiters=",;\t|")
    except csv.Error:
        dialect = csv.excel
    rows = list(csv.reader(io.StringIO(text), dialect))
    return Extracted(paras=table_to_paras(rows))


# ---------------------------------------------------------------------------
# Formati vecchi (.doc, .ppt, .wps...) tramite i programmi installati

def _via_libreoffice(path: Path, opts: Options, ctx: Context) -> Extracted:
    from .converter import extract  # import ritardato per evitare cicli

    with tempfile.TemporaryDirectory() as tmp:
        tmp_dir = Path(tmp)
        decrypted = decrypt_office(path, ctx)
        if decrypted is not None:
            path = tmp_dir / ("decifrato" + path.suffix)
            path.write_bytes(decrypted.read())
        target = {".ppt": "pptx", ".pps": "pptx", ".pot": "pptx", ".xls": "xlsx", ".xlt": "xlsx"}.get(path.suffix.lower(), "docx")
        converted = None
        try:
            converted = external.libreoffice_convert(path, target, tmp_dir)
        except Exception as exc:
            ctx.log(_("LibreOffice non è riuscito a convertire il file: {exc}").format(exc=exc))
        if converted is None and external.find_program("textutil") and target == "docx":
            out = tmp_dir / (path.stem + ".docx")
            external.run(["textutil", "-convert", "docx", "-output", str(out), str(path)])
            converted = out if out.is_file() else None
        if converted is None and target == "docx":
            converted = _word_com_convert(path, tmp_dir)
        if converted is None:
            raise ConversionError(_(
                "per i file {ext} serve LibreOffice (gratuito, www.libreoffice.org) "
                "oppure Microsoft Word installato sul computer.").format(ext=path.suffix))
        return extract(converted, opts, ctx)


def _word_com_convert(path: Path, out_dir: Path) -> Path | None:
    """Su Windows, se è installato Microsoft Word, lo usa per salvare in .docx."""
    import sys

    if sys.platform != "win32":
        return None
    try:
        import win32com.client  # type: ignore
    except ImportError:
        return None
    out = out_dir / (path.stem + ".docx")
    word = None
    try:
        word = win32com.client.DispatchEx("Word.Application")
        word.Visible = False
        doc = word.Documents.Open(str(path), ReadOnly=True, AddToRecentFiles=False)
        doc.SaveAs2(str(out), FileFormat=16)
        doc.Close(False)
    except Exception:
        return None
    finally:
        if word is not None:
            try:
                word.Quit()
            except Exception:
                pass
    return out if out.is_file() else None


def extract_legacy_office(path: Path, opts: Options, ctx: Context) -> Extracted:
    return _via_libreoffice(path, opts, ctx)
