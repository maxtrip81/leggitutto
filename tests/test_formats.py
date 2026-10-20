"""Prove di conversione su file creati al volo (non servono programmi esterni)."""

import zipfile

import pymupdf
import pytest

from leggitutto.converter import convert_file, extract
from leggitutto.model import Context, ProtectedError
from leggitutto.settings import Options

BODY = ("Il treno partì all'alba e attraversò la pianura ancora coperta di brina, mentre i passeggeri "
        "dormivano appoggiati ai finestrini e il controllore passava in silenzio fra i sedili. ")


def make_pdf(path, pages=4, password=None):
    doc = pymupdf.open()
    for n in range(1, pages + 1):
        page = doc.new_page(width=595, height=842)
        page.insert_text((72, 40), "Rivista di prova - numero speciale", fontsize=9)
        rect = pymupdf.Rect(72, 80, 523, 650)
        page.insert_textbox(rect, f"Paragrafo della pagina {n}. " + BODY * 3, fontsize=11, align=pymupdf.TEXT_ALIGN_JUSTIFY)
        page.insert_text((72, 720), f"1 Nota a piè di pagina numero {n}.", fontsize=8)
        page.insert_text((290, 800), str(n), fontsize=9)
    kwargs = {}
    if password:
        kwargs = dict(encryption=pymupdf.PDF_ENCRYPT_AES_256, user_pw=password, owner_pw=password + "x")
    doc.save(str(path), **kwargs)


def test_pdf_cleanup(tmp_path):
    src = tmp_path / "rivista.pdf"
    make_pdf(src)
    out = convert_file(src, Options(), Context())
    text = out.read_text(encoding="utf-8")
    assert "Rivista di prova" not in text
    assert "Nota a piè" not in text
    for n in range(1, 5):
        assert f"Paragrafo della pagina {n}." in text
    assert not any(line.strip().isdigit() for line in text.splitlines())


def test_pdf_with_password(tmp_path):
    src = tmp_path / "segreto.pdf"
    make_pdf(src, pages=1, password="ciao")
    with pytest.raises(ProtectedError):
        extract(src, Options(), Context())
    ctx = Context(ask_password=lambda name, attempt: "ciao")
    assert "Paragrafo della pagina 1." in convert_file(src, Options(), ctx).read_text(encoding="utf-8")


def test_docx_table_and_headings(tmp_path):
    import docx

    d = docx.Document()
    d.add_heading("Relazione", 1)
    d.add_paragraph("Primo paragrafo.")
    table = d.add_table(rows=2, cols=2)
    table.cell(0, 0).text, table.cell(0, 1).text = "Nome", "Ruolo"
    table.cell(1, 0).text, table.cell(1, 1).text = "Anna", "Direttrice"
    d.sections[0].header.paragraphs[0].text = "Intestazione da togliere"
    src = tmp_path / "relazione.docx"
    d.save(src)
    text = convert_file(src, Options(), Context()).read_text(encoding="utf-8")
    assert text.startswith("Relazione\n\nPrimo paragrafo.")
    assert "Nome: Anna; Ruolo: Direttrice." in text
    assert "Intestazione" not in text


def test_html_reader_mode_and_notes(tmp_path):
    src = tmp_path / "pagina.html"
    src.write_text("""<html><body><nav>Menu Home Contatti</nav><main>
        <h1>Titolo</h1><p>Testo con nota<sup><a href="#fn1">1</a></sup> nel mezzo.</p>
        <p class="footnote" id="fn1"><a href="#r1">1</a> Testo della nota.</p></main>
        <footer>Copyright</footer></body></html>""", encoding="utf-8")
    text = convert_file(src, Options(), Context()).read_text(encoding="utf-8")
    assert text == "Titolo\n\nTesto con nota nel mezzo.\n"


def test_epub_with_drm_is_refused(tmp_path):
    src = tmp_path / "libro.epub"
    with zipfile.ZipFile(src, "w") as zf:
        zf.writestr("mimetype", "application/epub+zip")
        zf.writestr("META-INF/rights.xml", "<rights/>")
        zf.writestr("META-INF/container.xml", "<container/>")
    with pytest.raises(ProtectedError):
        extract(src, Options(), Context())


def test_subtitles(tmp_path):
    src = tmp_path / "film.srt"
    src.write_text("1\n00:00:01,000 --> 00:00:02,000\nCiao\n\n2\n00:00:02,100 --> 00:00:03,000\n<i>a tutti</i>\n", encoding="utf-8")
    assert convert_file(src, Options(), Context()).read_text(encoding="utf-8") == "Ciao a tutti\n"


def test_unknown_extension_is_sniffed(tmp_path):
    src = tmp_path / "documento.bin"
    make_pdf(src, pages=1)
    assert "Paragrafo della pagina 1." in finalize_text(src)


def finalize_text(src):
    from leggitutto.cleaner import finalize

    return finalize(extract(src, Options(), Context()), Options())


def test_speaker_assignment():
    from leggitutto.media import Word, _paragraphs, assign_speakers

    words = [Word(i * 0.5, i * 0.5 + 0.4, f" w{i}") for i in range(12)]
    turns = [(0.0, 2.9, 0), (3.0, 3.4, 1), (3.5, 6.0, 1)]
    assign_speakers(words, turns)
    assert [w.speaker for w in words[:6]] == [0] * 6
    assert all(w.speaker == 1 for w in words[6:])
    paras = _paragraphs(words, True)
    assert [p[1] for p in paras] == [0, 1]
